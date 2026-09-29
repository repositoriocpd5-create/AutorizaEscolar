"""Consultas administrativas, indicadores e exportação (CSV, XLSX, PDF)."""
import csv
import io
import unicodedata
from dataclasses import dataclass
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func
from sqlalchemy.orm import selectinload

from ..extensions import db
from ..formatacao import agora_local, data_br, data_hora_utc_br, rotulo_situacao
from ..models import (Aluno, Autorizacao, Escola, Passeio, PasseioAluno, Responsavel, ResponsavelAluno, Situacao,
                      Turma)
from .passeio_service import estado_passeio


@dataclass
class Linha:
    aluno: Aluno
    autorizacao: Autorizacao | None
    situacao: str
    responsavel_nome: str
    responsaveis: str

    @property
    def data(self) -> str:
        return data_hora_utc_br(self.autorizacao.data_hora).replace(" às ", " ") if self.autorizacao else ""

    @property
    def protocolo(self) -> str:
        return self.autorizacao.protocolo or "" if self.autorizacao else ""


FILTROS = ("escola", "turma", "ano", "situacao", "q")

RELATORIOS = {
    "geral": ("Relação geral de alunos", {}),
    "autorizados": ("Relação de alunos autorizados", {"situacao": Situacao.AUTORIZADO}),
    "nao_autorizados": ("Relação de alunos não autorizados", {"situacao": Situacao.NAO_AUTORIZADO}),
    "aguardando": ("Responsáveis que ainda não responderam", {"situacao": Situacao.AGUARDANDO}),
    "por_turma": ("Lista por turma", {"_ordem": "turma"}),
    "por_escola": ("Lista por escola", {"_ordem": "escola"}),
}


def consultar(passeio: Passeio, filtros: dict, admin) -> list[Linha]:
    q = (Aluno.query.options(selectinload(Aluno.escola), selectinload(Aluno.turma))
         .join(PasseioAluno, PasseioAluno.aluno_id == Aluno.id)
         .join(Turma, Turma.id == Aluno.turma_id).join(Escola, Escola.id == Aluno.escola_id)
         .filter(PasseioAluno.passeio_id == passeio.id))
    if admin.escola_id:  # escopo do administrador
        q = q.filter(Aluno.escola_id == admin.escola_id)
    if filtros.get("escola"):
        q = q.filter(Escola.id == _int(filtros["escola"]))
    if filtros.get("turma"):
        q = q.filter(Turma.id == _int(filtros["turma"]))
    if filtros.get("ano"):
        q = q.filter(Turma.ano == filtros["ano"])
    alunos = q.order_by(Escola.nome, Turma.ano, Turma.nome, Aluno.nome).all()

    ids = [a.id for a in alunos] or [0]
    auts = {a.aluno_id: a for a in Autorizacao.query.options(
        selectinload(Autorizacao.responsavel), selectinload(Autorizacao.documento)).filter(
        Autorizacao.passeio_id == passeio.id, Autorizacao.aluno_id.in_(ids))}
    vinculos = (db.session.query(
                    ResponsavelAluno.aluno_id,
                    func.string_agg(Responsavel.nome, ", ").label("responsaveis"))
                .join(Responsavel, Responsavel.id == ResponsavelAluno.responsavel_id)
                .join(PasseioAluno, PasseioAluno.aluno_id == ResponsavelAluno.aluno_id)
                .filter(PasseioAluno.passeio_id == passeio.id, ResponsavelAluno.ativo.is_(True),
                        ResponsavelAluno.responsavel_legal.is_(True))
                .group_by(ResponsavelAluno.aluno_id))
    vinc = dict(vinculos)

    encerrado = estado_passeio(passeio).encerrado
    termo = _normalizar(filtros.get("q") or "")
    linhas = []
    for aluno in alunos:
        aut = auts.get(aluno.id)
        situacao = aut.situacao if aut else (Situacao.ENCERRADO if encerrado else Situacao.AGUARDANDO)
        nomes = vinc.get(aluno.id, "")
        resp_nome = aut.responsavel.nome if aut else (nomes.split(", ", 1)[0] if nomes else "—")
        linha = Linha(aluno, aut, situacao, resp_nome, nomes)
        if termo and not all(parte in _normalizar(" ".join(x or "" for x in (
                aluno.nome, aluno.matricula, linha.responsaveis, resp_nome, linha.protocolo)))
                for parte in termo.split()):
            continue
        filtro_sit = filtros.get("situacao")
        if filtro_sit:
            alvo = {Situacao.AGUARDANDO, Situacao.ENCERRADO} if filtro_sit == Situacao.AGUARDANDO else {filtro_sit}
            if situacao not in alvo:
                continue
        linhas.append(linha)
    if filtros.get("_ordem") == "turma":
        linhas.sort(key=lambda l: (l.aluno.turma.ano, l.aluno.turma.nome, l.aluno.escola.nome, l.aluno.nome))
    return linhas


def indicadores(linhas: list[Linha]) -> dict:
    total = len(linhas)
    cont = {s: 0 for s in Situacao.ROTULOS}
    for l in linhas:
        cont[l.situacao] += 1
    aguardando = cont[Situacao.AGUARDANDO] + cont[Situacao.ENCERRADO]
    pct = (cont[Situacao.AUTORIZADO] / total * 100) if total else 0.0
    return {"total": total, "autorizados": cont[Situacao.AUTORIZADO],
            "nao_autorizados": cont[Situacao.NAO_AUTORIZADO], "cancelados": cont[Situacao.CANCELADO],
            "aguardando": aguardando, "percentual": pct,
            "percentual_fmt": f"{pct:.1f}".replace(".", ",") + "%"}


def _normalizar(texto: str) -> str:
    """Minúsculas e sem acentos: 'José' encontra 'jose'."""
    decomposto = unicodedata.normalize("NFKD", texto.strip().lower())
    return "".join(c for c in decomposto if not unicodedata.combining(c))


def _int(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return -1


# ---------------------------------------------------------------------------
# Exportação — CPF nunca é exportado.
# ---------------------------------------------------------------------------
CABECALHO = ["Matrícula", "Aluno", "Escola", "Ano", "Turma", "Responsável", "Situação", "Data da resposta",
             "Protocolo"]


def _linhas_tabulares(linhas):
    for l in linhas:
        yield [l.aluno.matricula, l.aluno.nome, l.aluno.escola.nome, l.aluno.turma.ano, l.aluno.turma.nome,
               l.responsavel_nome, rotulo_situacao(l.situacao), l.data, l.protocolo]


def exportar_csv(linhas) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(CABECALHO)
    for row in _linhas_tabulares(linhas):
        w.writerow([_seguro_planilha(c) for c in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM p/ Excel pt-BR


def exportar_xlsx(linhas, titulo: str, passeio: Passeio) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"
    ws.append([titulo])
    ws.append([f"{passeio.nome} — {data_br(passeio.data)} — gerado em {agora_local():%d/%m/%Y %H:%M}"])
    ws.append([])
    ws.append(CABECALHO)
    ws["A1"].font = Font(bold=True, size=14)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0C1D32")
        cell.alignment = Alignment(vertical="center")
    for row in _linhas_tabulares(linhas):
        ws.append([_seguro_planilha(c) for c in row])
    for col, largura in zip("ABCDEFGHI", (12, 32, 30, 10, 10, 30, 22, 18, 18)):
        ws.column_dimensions[col].width = largura
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{ws.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar_pdf(linhas, titulo: str, passeio: Passeio, instituicao: str, brasao: str | None = None) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=titulo)
    h = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=14, leading=18, textColor=colors.HexColor("#007AAD"))
    n = ParagraphStyle("n", fontName="Helvetica", fontSize=9, leading=12)
    c = ParagraphStyle("c", fontName="Helvetica", fontSize=8, leading=10)
    ind = indicadores(linhas)
    # Reutiliza a mesma regra do documento de autorização: brasão configurado
    # no painel, com retorno ao brasão institucional padrão.
    from .documento_service import _brasao
    cabecalho = Table([[_brasao(brasao, largura=18 * mm), [
        Paragraph(instituicao.upper(), n),
        Paragraph(titulo, h),
        Paragraph(f"{passeio.nome} — {passeio.destino} — {data_br(passeio.data)}", n),
        Paragraph(f"Total: {ind['total']} · Autorizados: {ind['autorizados']} · Não autorizados: "
                  f"{ind['nao_autorizados']} · Aguardando: {ind['aguardando']} · Gerado em "
                  f"{agora_local():%d/%m/%Y %H:%M}", n),
    ]]], colWidths=[22 * mm, None])
    cabecalho.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    el = [cabecalho, Spacer(1, 6)]
    dados = [[Paragraph(f"<b>{x}</b>", c) for x in CABECALHO]]
    dados += [[Paragraph(escape(str(x)), c) for x in row] for row in _linhas_tabulares(linhas)]
    t = Table(dados, repeatRows=1, colWidths=[20 * mm, 50 * mm, 45 * mm, 17 * mm, 17 * mm, 45 * mm,
                                               30 * mm, 26 * mm, 25 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C1CE")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F7")),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    el.append(t)
    doc.build(el)
    return buf.getvalue()


def _seguro_planilha(valor):
    """Evita injeção de fórmulas em planilhas (CSV/XLSX injection)."""
    s = str(valor)
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s
